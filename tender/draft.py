"""Tool 5: draft_response. A first-draft bid response (.docx) for a human bid writer to finish.

    plan   one section per evaluation criterion, word budget from page limit × weight share   (code)
    write  one structured LLM call per section, only profile facts, placeholders where proof is missing
    check  word budget, unverified numbers / program codes, placeholder list                (code)
    build  .docx: cover marked AI DRAFT, compliance table, sections, yellow placeholders, notes (code)

Only runs after a person approved the bid (require_approved). Output goes to Blob: draft/draft.docx + draft.json.
"""
import io
import re
from datetime import datetime, timezone

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt
from pydantic import BaseModel, Field

from config import CHAT_DEPLOYMENT, aoai
from tender import store
from tender.tools import _norm, load_profile, require_approved

DRAFT_PROMPT_VERSION = "draft-v2"
MIN_BUDGET_SHARE = 0.6   # a section below 60% of its budget wastes evaluation marks
WORDS_PER_PAGE = 450
DEFAULT_RESPONSE_PAGES = 10
PLACEHOLDER = re.compile(r"\[(?:EVIDENCE NEEDED|CONFIRM):[^\]]+\]")
PRICE_CRITERION = re.compile(r"value for money|price|cost", re.IGNORECASE)

# ── plan (code) ───────────────────────────────────────────────────────────────


def _pages(limit: str | None) -> int | None:
    m = re.search(r"(\d+)\s*pages?", limit or "", re.IGNORECASE)
    return int(m.group(1)) if m else None


def plan_sections(req: dict) -> list[dict]:
    """One section per evaluation criterion. Word budget = criteria pages × words/page × weight share.

    Price criteria get a short pointer to the pricing schedule instead of an essay.
    """
    criteria_part = next((p for p in req.get("response_parts", [])
                          if re.search(r"evaluation criteria", p["content"], re.IGNORECASE)), None)
    pages = _pages(criteria_part["page_limit"]) if criteria_part else None
    total_words = (pages or DEFAULT_RESPONSE_PAGES) * WORDS_PER_PAGE
    written = [c for c in req["evaluation_criteria"] if not PRICE_CRITERION.search(c["name"])]
    weights = [c["weight_percent"] or 0 for c in written]
    total_weight = sum(weights) or len(written) or 1
    plan = []
    for c in req["evaluation_criteria"]:
        is_price = bool(PRICE_CRITERION.search(c["name"]))
        share = 0 if is_price else ((c["weight_percent"] or 0) / total_weight if sum(weights) else 1 / len(written))
        plan.append({"number": c["number"], "criterion": c["name"], "weight_percent": c["weight_percent"],
                     "price_criterion": is_price, "word_budget": 60 if is_price else max(120, int(total_words * share)),
                     "response_part": criteria_part["part"] if criteria_part else None})
    return plan


# ── write (LLM) ───────────────────────────────────────────────────────────────

SECTION_SYSTEM = """You are drafting one section of a government tender response for the company described in
the PROFILE. The draft will be finished by a human bid writer.

Rules:
- Address the CRITERION directly and persuasively, in the company's voice ("we").
- Use ONLY facts stated in the PROFILE and the TENDER. Never invent clients, projects, numbers, dates,
  certifications, staff or awards.
- Where a strong response needs evidence the PROFILE does not contain, insert a placeholder exactly like
  [EVIDENCE NEEDED: what is needed] and keep going.
- LENGTH: write between 70% and 100% of the word budget. Page allowance is evaluation opportunity; do not
  leave it unused.
- STRUCTURE the section as 4-6 paragraphs: (1) our understanding of what the criterion asks, (2) our approach,
  (3) relevant experience from the PROFILE, (4) how we will deliver for this tender's scope, (5) evidence or
  placeholders for evidence still needed.
- Plain, specific sentences tied to this tender's scope. No generic claims such as "best practice" or
  "proven methodology" unless the PROFILE supports them with specifics."""


class Section(BaseModel):
    heading: str
    paragraphs: list[str] = Field(description="The section body, 4-6 paragraphs")
    profile_facts_used: list[str] = Field(description="Short quotes from the PROFILE that this section relies on")


def write_section(tender: dict, plan: dict, req: dict, profile: str) -> tuple[dict, int]:
    scope = "\n".join(f"- {s}" for s in req.get("scope_items", []))
    user = (f"TENDER: {tender['title']} ({tender['agency']})\nSCOPE:\n{scope}\n\n"
            f"CRITERION {plan['number'] or ''}: {plan['criterion']} (weighting {plan['weight_percent']}%)\n"
            f"WORD BUDGET: about {plan['word_budget']} words\n\nPROFILE:\n{profile}")
    resp = aoai.chat.completions.parse(
        model=CHAT_DEPLOYMENT, temperature=0, response_format=Section,
        messages=[{"role": "system", "content": SECTION_SYSTEM}, {"role": "user", "content": user}])
    return resp.choices[0].message.parsed.model_dump(), resp.usage.total_tokens


def compliance_rows(req: dict, score: dict | None) -> list[dict]:
    """Each mandatory condition with how it is met, taken from score_fit's evidence; unknowns become [CONFIRM]."""
    checks = (score or {}).get("mandatory_checks", [])
    rows = []
    for c in req.get("conditions", []):
        match = next((k for k in checks if _norm(k["tender_quote"])[:40] in _norm(c["tender_quote"])
                      or _norm(c["tender_quote"])[:40] in _norm(k["tender_quote"])), None)
        if match and match["met"] == "yes" and match.get("evidence"):
            response = f"Met. {match['evidence']}"
        else:
            response = f"[CONFIRM: evidence that we meet this condition ({c['id'] or 'condition'})]"
        rows.append({"id": c["id"], "condition": c["text"], "response": response})
    return rows


# ── check (code) ──────────────────────────────────────────────────────────────

# Figures and program codes: "LAND 400", "12,000", "98%", "2026". (?<!\w)/(?!\w) instead of \b so "98%" matches:
# \b needs a word character on one side, and "%" followed by a space has none.
CLAIM = re.compile(r"(?<!\w)(?:[A-Z]{2,}[ -]?\d{2,}[A-Z0-9-]*|\d{1,3}(?:,\d{3})+|\d+(?:\.\d+)?%|\d{3,})(?!\w)")


def unverified_claims(text: str, sources: str) -> list[str]:
    """Numbers and program codes in the draft that appear nowhere in the profile or tender: possible inventions."""
    src = _norm(sources)
    found = {m.group(0) for m in CLAIM.finditer(PLACEHOLDER.sub(" ", text))}
    return sorted(c for c in found if _norm(c) not in src)


def check_section(section: dict, plan: dict, sources: str) -> list[str]:
    body = " ".join(section["paragraphs"])
    warnings = []
    words = len(body.split())
    if words > plan["word_budget"] * 1.15:
        warnings.append(f"{words} words, over the ~{plan['word_budget']}-word budget")
    elif not plan["price_criterion"] and words < plan["word_budget"] * MIN_BUDGET_SHARE:
        warnings.append(f"{words} words, under {MIN_BUDGET_SHARE:.0%} of the ~{plan['word_budget']}-word budget")
    claims = unverified_claims(body, sources)
    if claims:
        warnings.append(f"unverified figures or codes (not in profile or tender): {claims}")
    return warnings


# ── build (code) ──────────────────────────────────────────────────────────────

def _paragraph_with_highlights(doc, text: str, style: str | None = None):
    """Add a paragraph; placeholders are highlighted yellow so the human writer can't miss them."""
    p = doc.add_paragraph(style=style)
    pos = 0
    for m in PLACEHOLDER.finditer(text):
        if m.start() > pos:
            p.add_run(text[pos:m.start()])
        run = p.add_run(m.group(0))
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        run.bold = True
        pos = m.end()
    if pos < len(text):
        p.add_run(text[pos:])
    return p


def build_docx(draft: dict) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(draft["title"], 0)
    _paragraph_with_highlights(doc, f"Response to {draft['agency']} · Tender {draft['tender_id']}")
    banner = doc.add_paragraph()
    run = banner.add_run("AI DRAFT – for human review. Not for submission until every highlighted item is resolved.")
    run.bold = True
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    doc.add_heading("Compliance with conditions for participation", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("ID", "Condition", "Our response"), strict=True):
        cell.text = text
    for row in draft["compliance"]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = row["id"] or "", row["condition"]
        cells[2].paragraphs[0].text = ""
        cells[2].paragraphs[0].add_run(row["response"]).font.highlight_color = (
            WD_COLOR_INDEX.YELLOW if PLACEHOLDER.search(row["response"]) else None)

    for s in draft["sections"]:
        title = f"Criterion {s['number']}: {s['criterion']}" if s["number"] else s["criterion"]
        doc.add_heading(f"{title} ({s['weight_percent']:g}%)" if s["weight_percent"] else title, 1)
        for para in s["paragraphs"]:
            _paragraph_with_highlights(doc, para)

    doc.add_page_break()
    doc.add_heading("Draft notes (delete before submission)", 1)
    for line in draft["notes"]:
        doc.add_paragraph(line, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── orchestrate ───────────────────────────────────────────────────────────────

def draft_response(tender_id: str) -> dict:
    """AGENT tool. Builds the draft .docx; refuses unless a person approved the bid."""
    try:
        require_approved(tender_id)
    except PermissionError as e:
        return {"error": str(e)}
    req = store.get_json(tender_id, "requirements.json")
    if not req or not req.get("evaluation_criteria"):
        return {"error": "No evaluation criteria found. Run extract_requirements first (tender documents needed)."}
    notice = store.get_json(tender_id, "notice.json")
    score = store.get_json(tender_id, "score.json")
    profile = load_profile()
    sources = profile + "\n" + str(notice) + "\n" + str(req)

    sections, notes, tokens = [], [], 0
    for plan in plan_sections(req):
        if plan["price_criterion"]:
            sec = {"heading": plan["criterion"], "profile_facts_used": [], "paragraphs": [
                "Our pricing is provided in the pricing schedule. "
                "[CONFIRM: rates and value-for-money narrative from the commercial lead]"]}
        else:
            sec, used = write_section(notice, plan, req, profile)
            tokens += used
        warnings = check_section(sec, plan, sources)
        notes += [f"Criterion {plan['number']}: {w}" for w in warnings]
        sections.append({**plan, **sec, "warnings": warnings,
                         "words": len(" ".join(sec["paragraphs"]).split())})

    compliance = compliance_rows(req, score)
    all_text = " ".join(" ".join(s["paragraphs"]) for s in sections) + " " + " ".join(r["response"] for r in compliance)
    placeholders = sorted(set(PLACEHOLDER.findall(all_text)))
    notes += [f"Placeholder to resolve: {p}" for p in placeholders]
    if req.get("warnings"):
        notes += [f"Requirements extraction: {w}" for w in req["warnings"]]
    stamp = f"{datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC"
    notes.append(f"Generated {stamp} · {DRAFT_PROMPT_VERSION} · {CHAT_DEPLOYMENT}")

    draft = {"tender_id": tender_id, "title": notice["title"], "agency": notice["agency"],
             "compliance": compliance, "sections": sections, "placeholders": placeholders, "notes": notes,
             "prompt_version": DRAFT_PROMPT_VERSION, "tokens": tokens,
             "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds")}
    store.put_file(tender_id, "draft/draft.docx", build_docx(draft))
    store.put_json(tender_id, "draft/draft.json", draft)
    store.set_status(tender_id, "drafted", by="agent2.draft_response",
                     note=f"{len(sections)} sections, {len(placeholders)} placeholders")
    return {"tender_id": tender_id, "status": "drafted", "sections": len(sections),
            "words": sum(s["words"] for s in sections), "placeholders": placeholders,
            "warnings": [n for n in notes if not n.startswith(("Placeholder", "Generated"))],
            "files": ["draft/draft.docx", "draft/draft.json"], "tokens": tokens}
